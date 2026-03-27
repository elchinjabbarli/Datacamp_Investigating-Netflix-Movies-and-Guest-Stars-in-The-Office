import os
import logging
from pptx import Presentation
from pptx.util import Inches
from typing import Dict, Any, List, Optional

class PPTGenerator:
    """
    Handles PowerPoint generation with template support and defensive logic.
    """
    def __init__(self, template_path: Optional[str] = None):
        self.logger = logging.getLogger(__name__)
        self.prs = self._initialize_presentation(template_path)

    def _initialize_presentation(self, path: Optional[str]) -> Presentation:
        """
        Self-Correction: If template path is invalid or file is corrupted,
        fall back to a blank presentation instead of crashing.
        """
        if path and os.path.exists(path):
            try:
                return Presentation(path)
            except Exception as e:
                self.logger.error(f"Template yüklenemedi, varsayılan sunum kullanılıyor: {e}")
                return Presentation()
        return Presentation()

    def _get_safe_layout(self, hint: str):
        """
        Self-Correction: If requested layout doesn't exist, search for keywords
        or return the most common 'Title and Content' layout (index 1).
        """
        hint_lower = hint.lower()
        # Search for keyword matches in layout names
        for layout in self.prs.slide_layouts:
            if hint_lower in layout.name.lower():
                return layout

        # Fallback logic
        try:
            return self.prs.slide_layouts[1] # Usually 'Title and Content'
        except IndexError:
            return self.prs.slide_layouts[0] # Total fallback to first layout

    def create_presentation(self, structure: Dict[str, Any], image_paths: List[str], output_path: str) -> str:
        """
        Builds the presentation based on AI-generated structure.
        """
        try:
            slides_data = structure.get("slides", [])

            for slide_item in slides_data:
                layout_hint = slide_item.get("layout_hint", "Title and Content")
                layout = self._get_safe_layout(layout_hint)
                slide = self.prs.slides.add_slide(layout)

                # Title
                title_shape = slide.shapes.title
                if title_shape:
                    title_shape.text = slide_item.get("title", "Başlıksız Slayt")

                # Content
                self._add_text_content(slide, slide_item.get("content", []))

                # Image Handling
                img_idx = slide_item.get("image_index")
                if img_idx is not None and 0 <= img_idx < len(image_paths):
                    self._add_image_safely(slide, image_paths[img_idx])

            self.prs.save(output_path)
            return output_path
        except Exception as e:
            self.logger.error(f"PPTX üretim hatası: {e}")
            raise Exception(f"Sunum dosyası oluşturulamadı: {str(e)}")

    def _add_text_content(self, slide, content: Any):
        """
        Handles various content formats (string or list).
        """
        body_placeholder = None
        for shape in slide.placeholders:
            if shape.placeholder_format.type == 2: # Body
                body_placeholder = shape
                break

        if body_placeholder:
            tf = body_placeholder.text_frame
            if isinstance(content, list):
                for i, point in enumerate(content):
                    p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
                    p.text = str(point)
                    p.level = 0
            else:
                tf.text = str(content)

    def _add_image_safely(self, slide, img_path: str):
        """
        Self-Correction: Validate image exists, handle scaling to prevent
        images from covering the entire slide or going off-bounds.
        """
        if not os.path.exists(img_path):
            self.logger.warning(f"Görsel bulunamadı, atlanıyor: {img_path}")
            return

        # Attempt to find a picture placeholder first
        placeholder = None
        for shape in slide.placeholders:
            if shape.placeholder_format.type == 18: # Picture placeholder
                placeholder = shape
                break

        if placeholder:
            try:
                placeholder.insert_picture(img_path)
                return
            except Exception:
                pass # Fallback to manual placement if placeholder fails

        # Manual placement fallback (Right side, scaled)
        try:
            left = Inches(6)
            top = Inches(2)
            width = Inches(3.5)
            slide.shapes.add_picture(img_path, left, top, width=width)
        except Exception as e:
            self.logger.error(f"Resim ekleme hatası: {e}")
